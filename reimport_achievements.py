#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import sqlite3
import uuid
from datetime import datetime
import openpyxl

def parse_date(date_value):
    if date_value is None:
        return None
    if isinstance(date_value, datetime):
        return date_value
    if isinstance(date_value, str):
        try:
            return datetime.strptime(date_value, '%Y-%m-%d')
        except:
            try:
                return datetime.strptime(date_value, '%Y/%m/%d')
            except:
                return None
    return None

def import_achievements():
    conn = sqlite3.connect('data/achievement.db')
    cursor = conn.cursor()
    
    # 清空现有数据
    cursor.execute("DELETE FROM achievements")
    
    wb = openpyxl.load_workbook('成果全量249条.xlsx')
    ws = wb['产品成果明细登记薄']
    
    achievements_data = []
    
    for row in range(2, ws.max_row + 1):
        achievement_id = f"ACH_{str(uuid.uuid4())[:8].upper()}"
        
        # 按照Excel列顺序读取
        name = ws.cell(row=row, column=1).value  # A: 成果名称
        achievement_form = ws.cell(row=row, column=2).value  # B: 成果形态
        organization = ws.cell(row=row, column=3).value  # C: 所属机构
        department = ws.cell(row=row, column=4).value  # D: 部门
        version = ws.cell(row=row, column=5).value  # E: 成果版本
        has_baseline = ws.cell(row=row, column=6).value  # F: 是否有基线
        application_scenario = ws.cell(row=row, column=7).value  # G: 应用场景
        requirement_proposer = ws.cell(row=row, column=8).value  # H: 成果需求提出人
        achievement_target = ws.cell(row=row, column=9).value  # I: 成果目标描述
        function_list_file = ws.cell(row=row, column=10).value  # J: 成果对应功能清单
        acceptor = ws.cell(row=row, column=11).value  # K: 验收人（可多人）
        acceptance_method = ws.cell(row=row, column=12).value  # L: 成果验收方式
        planned_acceptance_date = parse_date(ws.cell(row=row, column=13).value)  # M: 计划验收日期
        actual_acceptance_date = parse_date(ws.cell(row=row, column=14).value)  # N: 实际验收日期
        demo_url = ws.cell(row=row, column=15).value  # O: DEMO地址
        deliverables = ws.cell(row=row, column=16).value  # P: 成果验收提交物
        code_repository_url = ws.cell(row=row, column=17).value  # Q: 代码仓库/在线文档地址
        status_str = ws.cell(row=row, column=18).value  # R: 成果状态
        sale_type = ws.cell(row=row, column=19).value  # S: 成果可售类型
        change_reason = ws.cell(row=row, column=20).value  # T: 异常变更原因
        product_id = ws.cell(row=row, column=21).value  # U: 关联产品
        related_project_name = ws.cell(row=row, column=22).value  # V: 关联项目
        related_order_name = ws.cell(row=row, column=23).value  # W: 关联订单
        related_order_id = ws.cell(row=row, column=24).value  # X: 关联订单编号
        package_ids = ws.cell(row=row, column=25).value  # Y: 关联套餐
        created_by = ws.cell(row=row, column=26).value  # Z: 创建人
        
        if not name:
            continue
        
        # 如果product_id为空，给一个默认值
        if not product_id:
            product_id = 'CP_UNKNOWN'
        
        # 转换状态
        status = 'PRE_REGISTER'
        pre_register_time = None
        register_time = None
        record_time = None
        
        if status_str:
            if '登记' in str(status_str):
                status = 'RECORDED'
                # 对于已登记成果，使用实际验收日期作为记录时间
                if actual_acceptance_date:
                    record_time = actual_acceptance_date
                else:
                    record_time = datetime.now()
            elif '注册' in str(status_str):
                status = 'REGISTER'
                # 对于已注册成果，使用创建时间作为注册时间
                register_time = datetime.now()
            else:
                status = 'PRE_REGISTER'
                pre_register_time = datetime.now()
        else:
            pre_register_time = datetime.now()
        
        # 转换是否有基线
        has_baseline_str = '有基线' if has_baseline == '有基线' else '无基线'
        
        achievements_data.append((
            achievement_id,
            product_id,
            None,  # product_name
            None,  # organization_id
            organization,
            None,  # department_id
            department,
            name,
            version,
            None,  # product_external_version
            has_baseline_str,
            requirement_proposer,
            achievement_form,
            sale_type,
            function_list_file,
            package_ids,
            application_scenario,
            None,  # module_id
            None,  # module_name
            None,  # type
            achievement_target,  # description
            acceptor,  # owner
            achievement_target,  # achievement_target
            planned_acceptance_date,
            acceptance_method,
            acceptor,
            None,  # related_project_id
            related_project_name,
            related_order_id,
            related_order_name,
            None,  # acceptance_requirements
            None,  # acceptance_organization
            change_reason,
            deliverables,
            code_repository_url,
            demo_url,
            actual_acceptance_date,
            status,
            None,  # estimated_acceptance_month
            pre_register_time,
            register_time,
            record_time,
            created_by,
            created_by,
            datetime.now(),
            datetime.now(),
            None,  # risk_tags
            None   # change_version
        ))
    
    cursor.executemany('''
        INSERT INTO achievements (
            id, product_id, product_name, organization_id, organization_name,
            department_id, department_name, name, version, product_external_version,
            has_baseline, requirement_proposer, achievement_form, sale_type, function_list_file,
            package_ids, application_scenario, module_id, module_name, type,
            description, owner, achievement_target, planned_acceptance_date, acceptance_method,
            acceptor, related_project_id, related_project_name, related_order_id, related_order_name,
            acceptance_requirements, acceptance_organization, change_reason, deliverables, code_repository_url,
            demo_url, actual_acceptance_date, status, estimated_acceptance_month, pre_register_time,
            register_time, record_time, created_by, updated_by, created_at,
            updated_at, risk_tags, change_version
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', achievements_data)
    
    conn.commit()
    print(f"成功导入 {len(achievements_data)} 条成果数据")
    
    cursor.execute('SELECT COUNT(*) FROM achievements')
    count = cursor.fetchone()[0]
    print(f"数据库中现有 {count} 条成果记录")
    
    # 验证前3条数据
    cursor.execute('SELECT name, version, organization_name, department_name, status FROM achievements LIMIT 3')
    print("\n前3条数据预览：")
    for i, row in enumerate(cursor.fetchall(), 1):
        print(f"{i}. {row[0]} - {row[1]} - {row[2]} - {row[3]} - {row[4]}")
    
    conn.close()

if __name__ == '__main__':
    import_achievements()
