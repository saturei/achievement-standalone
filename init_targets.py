#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import sqlite3
import uuid
from datetime import datetime
import openpyxl

def init_targets():
    conn = sqlite3.connect('database/achievement.db')
    cursor = conn.cursor()
    
    cursor.execute("DELETE FROM targets")
    
    wb = openpyxl.load_workbook('产品-xl.xlsx')
    ws = wb['产品（机构）']
    
    targets_data = []
    
    for row in range(3, 31):
        department = ws.cell(row=row, column=1).value
        organization = ws.cell(row=row, column=2).value
        category = ws.cell(row=row, column=3).value
        sub_category = ws.cell(row=row, column=4).value
        annual_target = ws.cell(row=row, column=5).value
        q1_target = ws.cell(row=row, column=6).value
        q2_target = ws.cell(row=row, column=7).value
        q3_target = ws.cell(row=row, column=8).value
        q4_target = ws.cell(row=row, column=9).value
        owner = ws.cell(row=row, column=10).value
        
        if not department or not organization:
            continue
        
        target_id = str(uuid.uuid4())
        
        target_type = sub_category
        if '签约收入' in str(sub_category):
            target_type = '签约收入'
        elif '确权收入' in str(sub_category):
            target_type = '确权收入'
        elif '研发成果' in str(sub_category):
            target_type = '研发成果'
        elif '费用' in str(sub_category):
            target_type = '费用'
        
        def to_decimal(value):
            if value is None:
                return 0.0
            if isinstance(value, (int, float)):
                return float(value)
            return 0.0
        
        targets_data.append((
            target_id,
            department,
            organization,
            category,
            sub_category,
            target_type,
            2026,
            to_decimal(annual_target),
            to_decimal(q1_target),
            to_decimal(q2_target),
            to_decimal(q3_target),
            to_decimal(q4_target),
            owner,
            datetime.now(),
            datetime.now()
        ))
    
    cursor.executemany('''
        INSERT INTO targets (
            id, department, organization, category, sub_category, target_type,
            year, annual_target, q1_target, q2_target, q3_target, q4_target,
            owner, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    ''', targets_data)
    
    conn.commit()
    print(f"成功导入 {len(targets_data)} 条目标数据")
    
    cursor.execute('''
        SELECT department, organization, category, sub_category, target_type, annual_target
        FROM targets
        LIMIT 5
    ''')
    
    print("\n前5条数据预览：")
    for row in cursor.fetchall():
        print(f"部门: {row[0]}, 机构: {row[1]}, 科目: {row[2]}, 细分目标: {row[3]}, 目标类型: {row[4]}, 年度目标: {row[5]}")
    
    conn.close()

if __name__ == '__main__':
    init_targets()
